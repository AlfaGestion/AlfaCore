"""Exporta el guion de demostración y su matriz de casos, únicamente con datos ficticios."""
import argparse
import json
import shutil
import zipfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def documentar(salida):
    m=json.loads((salida/'manifiesto_demo.json').read_text(encoding='utf-8'))
    verificacion=salida/'base/verificacion.json'
    sembrado=verificacion.exists() and json.loads(verificacion.read_text(encoding='utf-8'))[0].get('lote_sembrado')==1
    estado=f'Lote sembrado y verificado en {m["base_destino"]}. Conserva información histórica heredada.' if sembrado else 'Archivos generados. Siembra no verificada en esta carpeta.'
    manual=Path('src/AlfaCore/Docs/manual_demo_conciliacion_bancaria.md')
    shutil.copyfile(manual,salida/'Hoja_de_ruta_demostracion.md')
    wb=Workbook();inicio=wb.active;inicio.title='Inicio'
    for row in [ ['Demostración Conciliación y Cobranzas','Datos ficticios'],['Servidor',m['servidor_destino']],['Base',m['base_destino']],
        ['Estado',estado],['Período','01/10/2026 a 31/10/2026'],
        ['Clientes',len(m['clientes'])],['Comprobantes',len(m['documentos'])],['Renglones contables esperados',679],
        ['Movimientos únicos',len(m['movimientos'])],['Aviso',m['aviso']],['Guía','Hoja_de_ruta_demostracion.md'],
        ['Privacidad','Si se conservan datos anteriores de ALFANET, mostrar únicamente el lote DEMO.']]:inicio.append(row)
    ws=wb.create_sheet('Archivos');ws.append(['Archivo','Banco','Cuenta','Filas','Repetidos esperados','Nuevos esperados'])
    for f in m['archivos']:ws.append([f['archivo'],f['banco'],f['cuenta'],f['filas'],f['duplicados_esperados'],f['nuevos_esperados']])
    ws=wb.create_sheet('Casos');ws.append(['Banco','Cuenta','Caso','Objetivo','Fecha banco','Importe banco','Detalle ficticio','Referencia interna del guion','Comprobantes del caso','Resultado obtenido','Observaciones'])
    for mov in m['movimientos']:
        prefijo=int(next(i for i,b in enumerate(m['bancos'],1) if b['nombre']==mov['banco']))
        cuentas={c['codigo'] for c in m['clientes'] if int(c['codigo'][-3:])//100==prefijo}
        docs=[d['tc']+' '+d['comprobante']+' $ '+d['importe'] for d in m['documentos'] if d['caso']==mov['caso'] and d['cliente'] in cuentas]
        ws.append([mov['banco'],mov['cuenta'],mov['caso'],mov['objetivo'],mov['fecha'],float(mov['importe']),mov['descripcion'],mov['referencia'],' | '.join(docs),'',''])
    ws=wb.create_sheet('Documentos');ws.append(['Caso','Cliente','Tipo','Comprobante','Fecha','Importe','Banco','Retención','Aplicaciones'])
    for d in m['documentos']:ws.append([d['caso'],d['cliente'],d['tc'],d['comprobante'],d['fecha'],float(d['importe']),d['cuenta_banco'],float(d['retencion']),
        ' | '.join(a['tc']+' '+a['comprobante']+' $ '+a['importe'] for a in d['aplicaciones'])])
    ws=wb.create_sheet('Clientes');ws.append(['Código','Nombre ficticio','CUIT ficticio'])
    for c in m['clientes']:ws.append([c['codigo'],c['nombre'],c['cuit']])
    ws=wb.create_sheet('Checklist');ws.append(['Paso','Acción','Resultado esperado','Fecha','Presentador','Banco/cuenta','Resultado obtenido','Observaciones'])
    pasos=[('Preparar','Confirmar base y siembra','ALFANET / SERVER-ALFAVB6; carga definitiva completada'),
      ('Privacidad','Revisar alcance de la pantalla compartida','Solo datos DEMO; base exclusiva si presentación pública'),
      ('Importar','Importar archivo mensual de cada banco','61 movimientos por banco'),
      ('Buscar','Buscar importe en todos los bancos','Localizar la fila correcta'),('Ordenar','Pulsar encabezados y cambiar página','Orden y páginas conservan filtros'),
      ('E01','Confirmar propuesta exacta','Un vínculo; historial disponible'),('E02','Aceptar diferencia con motivo','0,22 dentro de tolerancia; importes originales intactos'),
      ('E03/E04','Revisar margen de fechas','Dos días dentro; siete días requieren revisión'),('E10','Abrir cuenta corriente y asiento completo','Recibo anterior visible; no crear otro'),
      ('E05','Cobrar dos proformas juntas','Un recibo, dos aplicaciones; saldo cero'),('E06','Revisar retención sugerida','Cuenta 11991 e importe restante 500,00'),
      ('E07','Cobrar parcialmente','Queda la mitad del saldo pendiente'),('E08','Asignar cliente y recordar','Asignación aplicada y propuesta reevaluada'),
      ('E11','Comparar clientes del mismo CUIT','Importe y evidencia distinguen las empresas'),('E13','Generar asiento de comisión','Asiento balanceado y conciliación trazable'),
      ('E17','Conciliar dos depósitos con un recibo','Totales iguales sin reutilizar renglones'),('E18','Revisar propuestas ambiguas','Sin confirmación automática indebida'),
      ('IA','Revisar propuestas con IA','Evidencia revisada; solo confirmar lo correcto'),('Ayuda','Preguntar a IA y abrir su fuente','Se abre el manual, se conserva la operación'),
      ('Duplicados','Importar segunda descarga','42 repetidos y 2 nuevos por cuenta; total 63'),('Reimportar','Repetir un mismo archivo','Cero movimientos nuevos'),
      ('Cerrar','Registrar resultados y límites de la demo','Checklist completo; sin prometer pasos no comprobados')]
    for paso,accion,esperado in pasos:ws.append([paso,accion,esperado,'','','','',''])
    for ws in wb.worksheets:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='264866')
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=28
        ws.row_dimensions[1].height=34
        for row in ws:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=46
        ws.sheet_view.zoomScale=85
    inicio.column_dimensions['A'].width=40;inicio.column_dimensions['B'].width=105
    for hoja,anchos in {'Casos':{'A':18,'B':16,'C':12,'D':44,'E':15,'F':18,'G':70,'H':24,'I':76,'J':42,'K':48},
        'Checklist':{'A':18,'B':48,'C':74,'D':17,'E':24,'F':26,'G':46,'H':50},'Documentos':{'D':24,'F':18,'I':76}}.items():
        for col,ancho in anchos.items():wb[hoja].column_dimensions[col].width=ancho
    wb['Casos'].freeze_panes='D2'
    for row in wb['Casos'].iter_rows(min_row=2,min_col=6,max_col=6):row[0].number_format='#,##0.00'
    wb.properties.creator='AlfaCore DEMO';wb.properties.lastModifiedBy='AlfaCore DEMO'
    wb.save(salida/'Guion_demostracion.xlsx')
    with zipfile.ZipFile(salida/'Extractos_DEMO_Conciliacion.zip','w',zipfile.ZIP_DEFLATED) as z:
        for archivo in [f['archivo'] for f in m['archivos']]+['manifiesto_demo.json','Guion_demostracion.xlsx','Hoja_de_ruta_demostracion.md']:
            z.write(salida/archivo,archivo)
    print('Guion Excel, hoja de ruta y paquete público generados. El respaldo original no se incluye en el ZIP.')

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--salida',type=Path,required=True)
    documentar(p.parse_args().salida)
